import os
import re
from difflib import SequenceMatcher
from .excel_extractor import AOC4ExcelExtractor

class AOC4PreviousYearReconciler:
    """
    Dedicated Additional Feature for AOC-4:
    Performs full line-by-line cross-reconciliation across ALL Financial Statement tables
    (Balance Sheet, Statement of Profit & Loss, and Notes to Accounts).
    Compares CY Financials (Previous Year Comparative Column) vs Last Year Filed Financials (Actual Filed Column).
    """
    def __init__(self, config_path: str = None):
        self.excel_extractor = AOC4ExcelExtractor()

    def _clean_number(self, val_str: str):
        """Extracts float from currency string, handling parentheses, dashes, commas."""
        if not val_str or not isinstance(val_str, str):
            return None
        clean = val_str.strip().replace(",", "").replace(" ", "")
        if clean in ["-", "–", "—", "nil", "NIL", "Nil", "pt", "_", ""]:
            return 0.0
        # Check for percentage
        if "%" in clean:
            return None
        # Handle parentheses for negative numbers (e.g. (17,719.08) -> -17719.08)
        if clean.startswith("(") and clean.endswith(")"):
            clean = "-" + clean[1:-1]
        try:
            return float(clean)
        except ValueError:
            return None

    def _clean_spelling(self, title: str) -> str:
        """Corrects OCR drop-letter typos and formatting noise in financial line item titles."""
        t = title.strip()
        typos = {
            r'^fotal\b': 'Total',
            r'^rade receivables': 'Trade receivables',
            r'^ash and cash equivalents': 'Cash and cash equivalents',
            r'^hort-term loans and advances': 'Short-term loans and advances',
            r'^ther current assets': 'Other current assets',
            r'^urrent assets': 'Current assets',
            r'^re capital at the beginning of the period': 'Share capital at the beginning of the period',
            r'^apital at the end of the period': 'Share capital at the end of the period',
            r'^eave encashment': 'Leave Encashment',
            r'^epreciation of property': 'Depreciation of Property, Plant and Equipment',
            r'^mortization of intangible': 'Amortization of Intangible assets',
            r'^npairment of goodwill': 'Impairment of goodwill',
            r'^rofessional charges': 'Professional Charges',
            r'^ncility costs.*': 'Facility Costs, Repairs & Maintenance',
            r'^ates and taxes': 'Rates and Taxes',
            r'^ecruitment and hiring charges': 'Recruitment and hiring charges',
            r'^elephone & internet charges': 'Telephone & Internet Charges',
            r'^fice expenses': 'Office Expenses',
            r'^гах expense': 'Tax expense',
            r'^arnings per equity share': 'Earnings per equity share',
            r'^ignificant accounting policies': 'Significant accounting policies'
        }
        for pat, rep in typos.items():
            if re.search(pat, t, re.IGNORECASE):
                t = re.sub(pat, rep, t, flags=re.IGNORECASE)
                break
        t = t.replace('\n', ' ').replace('**', '').replace('__', '').strip()
        return t

    def _normalize_title(self, title: str) -> str:
        """Normalizes line item title for robust matching."""
        clean = self._clean_spelling(title)
        t = clean.lower()
        t = re.sub(r'^\s*\(?[a-z0-9ivx]+\)?[.)\s]*', '', t) # remove numbering like (a), 1., etc.
        t = re.sub(r'[^a-z0-9 ]', ' ', t)
        t = re.sub(r'\s+', ' ', t).strip()
        return t

    def _extract_table_rows(self, text: str, scale: float = 1.0) -> list:
        """
        Extracts all financial statement table rows containing particulars and figures
        with table-header column awareness and OCR spelling corrections.
        """
        lines = text.split("\n")
        tables = []
        curr_table = []
        curr_section = "Financial Statements"
        in_board_report = False
        
        for idx, line in enumerate(lines):
            line_clean = line.strip()
            # Strip markdown formatting like #, *, _, <b>, </b> for clean heading classification
            norm_heading = re.sub(r'[*_#<>]|</?b>', '', line_clean).strip().lower()
            
            # Track and exclude Board Report sections from Financial Statement reconciliations
            if re.search(r'^(board(\'s)?\s*report|directors?\'?\s*report)', norm_heading):
                in_board_report = True
                continue
            elif re.search(r'^(independent\s*auditor\'s\s*report|auditor\'s\s*report|notes\s*to\s*(the\s*)?financial|notes\s*to\s*accounts|balance\s*sheet|statement\s*of\s*profit)', norm_heading):
                in_board_report = False
                
            if in_board_report:
                continue
                
            if "balance sheet" in norm_heading:
                curr_section = "Balance Sheet"
            elif "statement of profit" in norm_heading or "profit and loss" in norm_heading:
                curr_section = "Statement of Profit and Loss"
            elif "cash flow" in norm_heading:
                curr_section = "Cash Flow Statement"
            elif norm_heading.startswith("note") or "notes to" in norm_heading:
                curr_section = f"Notes ({line_clean[:40]})"
                
            if line_clean.startswith("|") and line_clean.endswith("|"):
                curr_table.append((curr_section, line_clean))
            else:
                if len(curr_table) > 1:
                    tables.append(curr_table)
                curr_table = []
        if len(curr_table) > 1:
            tables.append(curr_table)
            
        extracted_rows = []
        for tbl in tables:
            sec = tbl[0][0]
            header_rows = []
            data_rows = []
            for s, row_str in tbl:
                cells = [c.strip() for c in row_str.strip("|").split("|")]
                # Skip markdown separator rows like |---|---|
                if all(re.match(r'^[-:\s]+$', c) for c in cells if c):
                    continue
                if not data_rows and any(h in cells[0].lower() for h in ["particulars", "description", "related party", "name of", "statement of", "balance sheet as at", "for the year ended"]):
                    header_rows.append(cells)
                else:
                    data_rows.append(cells)
                    
            # Determine column positions from header
            cy_col = None
            py_col = None
            part_cols = [0]
            
            if header_rows:
                all_hdr_text = " ".join([" ".join(h) for h in header_rows]).lower()
                years = sorted(list(set(int(y) for y in re.findall(r'\b(202\d)\b', all_hdr_text))), reverse=True)
                cy_yr = str(years[0]) if len(years) >= 1 else None
                py_yr = str(years[1]) if len(years) >= 2 else None
                
                if cy_yr:
                    cur_cy = int(self.detected_cy_year) if getattr(self, 'detected_cy_year', None) else 0
                    # Prioritize statutory BS / P&L headings or higher years
                    if int(cy_yr) > cur_cy or sec in ["Balance Sheet", "Statement of Profit and Loss"]:
                        self.detected_cy_year = cy_yr
                        if py_yr:
                            self.detected_py_year = py_yr
                        elif not getattr(self, 'detected_py_year', None):
                            self.detected_py_year = str(int(cy_yr) - 1)
                elif py_yr and not getattr(self, 'detected_py_year', None):
                    self.detected_py_year = py_yr

                # Find cy_col and py_col across all header rows
                for hdr in header_rows:
                    hdr_lower = [c.lower() for c in hdr]
                    for c_idx, c_name in enumerate(hdr_lower):
                        if cy_yr and cy_yr in c_name:
                            if cy_col is None: cy_col = c_idx
                        elif any(cy_kw in c_name for cy_kw in ["current year", "as at 31st march 202", "for the year ended 31st march 202"]):
                            if cy_col is None: cy_col = c_idx
                        
                        if py_yr and py_yr in c_name and c_idx != cy_col:
                            if py_col is None: py_col = c_idx
                        elif any(py_kw in c_name for py_kw in ["previous year", "comparative"]):
                            if py_col is None and c_idx != cy_col: py_col = c_idx

                # If "nature" or "description" is a secondary particulars column
                last_hdr_lower = [c.lower() for c in header_rows[-1]]
                if "nature" in last_hdr_lower:
                    part_cols = [0, last_hdr_lower.index("nature")]
                    
            for cells in data_rows:
                if len(cells) < 2:
                    continue
                raw_title = cells[0].strip()
                if not raw_title or len(raw_title) < 2:
                    continue
                if re.match(r'^(as at|for the year|inr|amount in|\d+&?\d*)$', raw_title.lower()):
                    continue
                    
                if len(part_cols) > 1 and len(cells) > part_cols[1] and cells[part_cols[1]]:
                    raw_title = f"{raw_title} - {cells[part_cols[1]].strip()}"
                    
                clean_title = self._clean_spelling(raw_title)
                
                # Determine CY and PY values based on identified columns or cell count
                cy_val = None
                py_val = None
                
                if cy_col is not None and len(cells) > cy_col and py_col is not None and len(cells) > py_col:
                    cy_val = self._clean_number(cells[cy_col])
                    py_val = self._clean_number(cells[py_col])
                else:
                    # Fallback: inspect numerical cells
                    num_cells = []
                    for c_idx, c in enumerate(cells[1:], 1):
                        num = self._clean_number(c)
                        if num is not None:
                            # Skip standalone small integer note numbers (e.g. Note 3, 15)
                            if c_idx == 1 and len(cells) >= 4 and num.is_integer() and 1 <= num <= 50 and "." not in c:
                                continue
                            num_cells.append(num)
                    # If 4 numbers (e.g. Shares CY, Amount CY, Shares PY, Amount PY), take Amount CY (index 1) and Amount PY (index 3)
                    if len(num_cells) == 4:
                        cy_val = num_cells[1]
                        py_val = num_cells[3]
                    elif len(num_cells) >= 2:
                        cy_val = num_cells[0]
                        py_val = num_cells[1]
                    elif len(num_cells) == 1:
                        cy_val = num_cells[0]
                        
                if cy_val is not None or py_val is not None:
                    # Non-currency items (share count, EPS, face value, ratios) must not be scaled
                    is_non_currency = any(kw in clean_title.lower() for kw in ["shares", "number of shares", "face value", "eps", "per share", "par value", "weighted average", "ratio"])
                    item_scale = 1.0 if is_non_currency else scale
                    extracted_rows.append({
                        "section": sec,
                        "particulars": clean_title,
                        "norm_title": self._normalize_title(clean_title),
                        "cy_val": cy_val * item_scale if cy_val is not None else None,
                        "py_val": py_val * item_scale if py_val is not None else None
                    })
                    
        return extracted_rows

    def _extract_excel_rows(self, excel_path: str) -> list:
        """Extracts financial line items and comparative columns from an Excel workbook across all financial sheets."""
        import openpyxl
        excel_rows = []
        try:
            wb = openpyxl.load_workbook(excel_path, data_only=True)
            
            # Detect unit scale (e.g. Lakhs = 100,000, Crores = 10,000,000)
            full_header_text = []
            for sheet in wb.sheetnames:
                ws = wb[sheet]
                for r in range(1, min(15, ws.max_row + 1)):
                    for c in range(1, min(15, ws.max_column + 1)):
                        v = ws.cell(row=r, column=c).value
                        if v: full_header_text.append(str(v))
            scale = self.excel_extractor.detect_financials_scale(" ".join(full_header_text))
            
            for sheet in wb.sheetnames:
                s_lower = sheet.lower()
                # Match standard and abbreviated sheet names
                if not any(k in s_lower for k in ['bs', 'balance', 'pl', 'profit', 'note', 'sch', 'statement', 'add notes']):
                    continue
                if any(skip in s_lower for skip in ['master', 'audit entry', 'tb ', 'it dep', 'deffered']):
                    continue
                
                ws = wb[sheet]
                cy_col, py_col, part_col = None, None, 1
                header_row = 1
                
                # Scan first 15 rows for headers
                for r in range(1, min(15, ws.max_row + 1)):
                    for c in range(1, min(ws.max_column + 1, 25)):
                        val = str(ws.cell(row=r, column=c).value or '').lower()
                        if any(y in val for y in ['2026', '25-26', 'current year', '31.03.2026', '31/03/2026', 'as at 31st march, 2026', '31.03.26']):
                            cy_col = c
                            header_row = r
                        elif any(y in val for y in ['2025', '24-25', 'previous year', '31.03.2025', '31/03/2025', 'as at 31st march, 2025', '31.03.25']):
                            py_col = c
                            header_row = r
                        elif any(p in val for p in ['particular', 'description', 'line item', 'nature']):
                            part_col = c
                
                # If no explicit year header found, look for generic amount columns
                if not cy_col and not py_col:
                    for r in range(1, min(15, ws.max_row + 1)):
                        for c in range(1, min(ws.max_column + 1, 20)):
                            val = str(ws.cell(row=r, column=c).value or '').lower()
                            if any(a in val for a in ['amount', 'rs.', 'inr', 'rupees']):
                                if not cy_col:
                                    cy_col = c
                                    header_row = r
                                elif not py_col:
                                    py_col = c
                
                if not cy_col and not py_col:
                    continue
                    
                for r in range(header_row + 1, ws.max_row + 1):
                    p = ws.cell(row=r, column=part_col).value
                    # If part_col empty, try column 1 or 2
                    if not p or not str(p).strip():
                        p = ws.cell(row=r, column=1).value or ws.cell(row=r, column=2).value
                    if not p or not str(p).strip():
                        continue
                        
                    p_str = str(p).strip()
                    if any(h in p_str.lower() for h in ['particulars', 'membership', 'din:', 'partner', 'director', 'notes forming', 'for the year ended', 'as at 31']):
                        continue
                        
                    cy_v = ws.cell(row=r, column=cy_col).value if cy_col else None
                    py_v = ws.cell(row=r, column=py_col).value if py_col else None
                    
                    try:
                        cy_val = float(cy_v) if cy_v is not None and str(cy_v).strip() not in ['', '-', 'None'] else None
                    except (ValueError, TypeError):
                        cy_val = None
                    try:
                        py_val = float(py_v) if py_v is not None and str(py_v).strip() not in ['', '-', 'None'] else None
                    except (ValueError, TypeError):
                        py_val = None
                        
                    if cy_val is not None or py_val is not None:
                        is_non_currency = any(kw in p_str.lower() for kw in ["shares", "number of shares", "face value", "eps", "per share", "par value", "weighted average", "ratio"])
                        item_scale = 1.0 if is_non_currency else scale
                        excel_rows.append({
                            'section': sheet,
                            'particulars': p_str,
                            'norm_title': self._normalize_title(p_str),
                            'cy_val': cy_val * item_scale if cy_val is not None else None,
                            'py_val': py_val * item_scale if py_val is not None else None
                        })
        except Exception as e:
            print(f"[!] Error reading Excel rows from {excel_path}: {e}")
        return excel_rows

    def reconcile(self, cy_text_or_data, py_text_or_data) -> dict:
        """
        Performs full line-by-line cross-reconciliation across ALL Financial Statement tables.
        """
        self.detected_cy_year = None
        self.detected_py_year = None

        if isinstance(cy_text_or_data, str) and (cy_text_or_data.endswith('.xlsx') or cy_text_or_data.endswith('.xls')) and os.path.exists(cy_text_or_data):
            cy_rows = self._extract_excel_rows(cy_text_or_data)
        elif isinstance(cy_text_or_data, str):
            scale_cy = self.excel_extractor.detect_financials_scale(cy_text_or_data)
            cy_rows = self._extract_table_rows(cy_text_or_data, scale=scale_cy)
        else:
            cy_rows = []

        doc1_cy_year = self.detected_cy_year
        doc1_py_year = self.detected_py_year

        self.detected_cy_year = None
        self.detected_py_year = None

        if isinstance(py_text_or_data, str) and (py_text_or_data.endswith('.xlsx') or py_text_or_data.endswith('.xls')) and os.path.exists(py_text_or_data):
            py_rows = self._extract_excel_rows(py_text_or_data)
        elif isinstance(py_text_or_data, str):
            scale_py = self.excel_extractor.detect_financials_scale(py_text_or_data)
            py_rows = self._extract_table_rows(py_text_or_data, scale=scale_py)
        else:
            py_rows = py_text_or_data

        doc2_cy_year = self.detected_cy_year
        doc2_py_year = self.detected_py_year

        # Overlapping reconciled year:
        # We compare Doc 1's PY comparative column (e.g. 2025) against Doc 2's CY actual filed column (e.g. 2025)
        reconciled_year = doc1_py_year or doc2_cy_year or (str(int(doc1_cy_year) - 1) if doc1_cy_year else "2025")

        reconciliation_items = []
        matched_count = 0
        mismatches_count = 0
        seen_keys = set()

        for cy_row in cy_rows:
            if cy_row.get("py_val") is None:
                continue

            item_key = f"{cy_row['section']}_{cy_row['norm_title']}"
            if item_key in seen_keys:
                continue

            # Look for best match in PY rows
            best_match = None
            best_score = 0.0

            for py_row in py_rows:
                # In cross-document matching, check cy_val or py_val (whichever is available in the target document)
                if cy_text_or_data is py_text_or_data or cy_text_or_data == py_text_or_data:
                    val_to_check = py_row.get("py_val")
                else:
                    val_to_check = py_row.get("cy_val") if py_row.get("cy_val") is not None else py_row.get("py_val")
                if val_to_check is None:
                    continue

                if cy_row["norm_title"] == py_row["norm_title"]:
                    score = 1.0
                elif cy_row["norm_title"] in py_row["norm_title"] or py_row["norm_title"] in cy_row["norm_title"]:
                    score = 0.85
                else:
                    score = SequenceMatcher(None, cy_row["norm_title"], py_row["norm_title"]).ratio()

                if score > 0.60 and score > best_score:
                    best_score = score
                    best_match = py_row

            if best_match:
                seen_keys.add(item_key)

                v_cy_py = cy_row["py_val"]
                if cy_text_or_data is py_text_or_data or cy_text_or_data == py_text_or_data:
                    v_py_actual = best_match["py_val"]
                else:
                    v_py_actual = best_match.get("cy_val") if best_match.get("cy_val") is not None else best_match.get("py_val")

                diff = round(v_cy_py - v_py_actual, 2)

                # Tolerance: within 1 unit or 0.1% for minor rounding differences
                is_match = abs(diff) <= max(1.0, abs(v_py_actual) * 0.001)
                status = "MATCHED" if is_match else "MISMATCH"

                if is_match:
                    matched_count += 1
                else:
                    mismatches_count += 1

                reconciliation_items.append({
                    "section": cy_row["section"],
                    "line_item": cy_row["particulars"],
                    "cy_reported_py": v_cy_py,
                    "last_year_filed": v_py_actual,
                    "variance": diff,
                    "status": status
                })

        return {
            "summary": {
                "total_compared": len(reconciliation_items),
                "matched": matched_count,
                "mismatches": mismatches_count,
                "status": "ALL_MATCHED" if mismatches_count == 0 and len(reconciliation_items) > 0 else ("MISMATCH_FOUND" if mismatches_count > 0 else "NO_DATA"),
                "cy_year": doc1_cy_year or "2026",
                "py_year": reconciled_year,
                "reconciled_year": reconciled_year
            },
            "reconciliation_items": reconciliation_items
        }

    def export_to_excel(self, report_result: dict, output_path: str):
        """Exports the full line-by-line variance report to a standalone Excel spreadsheet."""
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Full Statement PY Variance"

        # Title Block
        ws.merge_cells("A1:F1")
        ws["A1"] = "PREVIOUS YEAR FIGURES COMPREHENSIVE RECONCILIATION REPORT"
        ws["A1"].font = Font(name="Arial", size=13, bold=True, color="FFFFFF")
        ws["A1"].fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
        ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 28

        # Summary Block
        sum_data = report_result.get("summary", {})
        ws["A2"] = f"Total Line Items Reconciled: {sum_data.get('total_compared', 0)} | Matched: {sum_data.get('matched', 0)} | Mismatches: {sum_data.get('mismatches', 0)}"
        ws["A2"].font = Font(name="Arial", size=10, italic=True)
        ws.merge_cells("A2:F2")

        # Headers
        cy_str = sum_data.get("cy_year")
        py_str = sum_data.get("py_year")
        
        col_3_name = f"CY FS (FY {py_str} Comparative Col)" if py_str else "CY FS (PY Comparative Col)"
        col_4_name = f"Last Year Filed Financials (FY {py_str})" if py_str else "Last Year Filed Financials"

        headers = ["S.No", "Financial Section", "Financial Line Item", col_3_name, col_4_name, "Variance / Difference", "Status"]
        ws.append([]) # Blank row 3
        ws.append(headers) # Row 4

        header_font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="2F5597", end_color="2F5597", fill_type="solid")

        for col_idx, h in enumerate(headers, 1):
            cell = ws.cell(row=4, column=col_idx)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[4].height = 22

        # Populate rows
        items = report_result.get("reconciliation_items", [])
        for idx, item in enumerate(items, 1):
            status = item["status"]
            status_fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid") if status == "MATCHED" else PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")
            status_font = Font(name="Arial", size=9, bold=True, color="375623" if status == "MATCHED" else "C65911")

            row_cells = [
                idx,
                item.get("section", ""),
                item["line_item"],
                item["cy_reported_py"],
                item["last_year_filed"],
                item["variance"],
                status
            ]
            ws.append(row_cells)
            curr_row = ws.max_row
            
            # Format numbers
            ws.cell(row=curr_row, column=4).number_format = '#,##0.00'
            ws.cell(row=curr_row, column=5).number_format = '#,##0.00'
            ws.cell(row=curr_row, column=6).number_format = '#,##0.00'
            
            # Format status cell
            stat_cell = ws.cell(row=curr_row, column=7)
            stat_cell.fill = status_fill
            stat_cell.font = status_font
            stat_cell.alignment = Alignment(horizontal="center")

        # Auto-adjust column widths
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 14)

        os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
        wb.save(output_path)
        print(f"[+] Saved Comprehensive PY Variance Report to: {output_path}")
        return output_path

    def append_to_existing_workbook(self, report_result: dict, workbook_path: str, sheet_name: str = "Previous Year Comparison"):
        """Appends the formatted line-by-line reconciliation sheet into an existing populated AOC-4 workbook."""
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

        if not os.path.exists(workbook_path):
            return self.export_to_excel(report_result, workbook_path)

        wb = openpyxl.load_workbook(workbook_path)
        if sheet_name in wb.sheetnames:
            del wb[sheet_name]

        ws = wb.create_sheet(title=sheet_name)

        # Title Block
        ws.merge_cells("A1:G1")
        ws["A1"] = "PREVIOUS YEAR FIGURES COMPREHENSIVE RECONCILIATION REPORT"
        ws["A1"].font = Font(name="Arial", size=13, bold=True, color="FFFFFF")
        ws["A1"].fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
        ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 28

        # Summary Block
        sum_data = report_result.get("summary", {})
        ws["A2"] = f"Total Line Items Reconciled: {sum_data.get('total_compared', 0)} | Matched: {sum_data.get('matched', 0)} | Mismatches: {sum_data.get('mismatches', 0)}"
        ws["A2"].font = Font(name="Arial", size=10, italic=True)
        ws.merge_cells("A2:G2")

        # Headers
        cy_str = sum_data.get("cy_year")
        py_str = sum_data.get("py_year")
        
        col_3_name = f"CY FS (FY {py_str} Comparative Col)" if py_str else "CY FS (PY Comparative Col)"
        col_4_name = f"Last Year Filed Financials (FY {py_str})" if py_str else "Last Year Filed Financials"

        headers = ["S.No", "Financial Section", "Financial Line Item", col_3_name, col_4_name, "Variance / Difference", "Status"]
        ws.append([]) # Blank row 3
        ws.append(headers) # Row 4

        header_font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="2F5597", end_color="2F5597", fill_type="solid")

        for col_idx, h in enumerate(headers, 1):
            cell = ws.cell(row=4, column=col_idx)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[4].height = 22

        # Populate rows
        items = report_result.get("reconciliation_items", [])
        for idx, item in enumerate(items, 1):
            status = item["status"]
            status_fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid") if status == "MATCHED" else PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")
            status_font = Font(name="Arial", size=9, bold=True, color="375623" if status == "MATCHED" else "C65911")

            row_cells = [
                idx,
                item.get("section", ""),
                item["line_item"],
                item["cy_reported_py"],
                item["last_year_filed"],
                item["variance"],
                status
            ]
            ws.append(row_cells)
            curr_row = ws.max_row
            
            # Format numbers
            ws.cell(row=curr_row, column=4).number_format = '#,##0.00'
            ws.cell(row=curr_row, column=5).number_format = '#,##0.00'
            ws.cell(row=curr_row, column=6).number_format = '#,##0.00'
            
            # Format status cell
            stat_cell = ws.cell(row=curr_row, column=7)
            stat_cell.fill = status_fill
            stat_cell.font = status_font
            stat_cell.alignment = Alignment(horizontal="center")

        # Auto-adjust column widths
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 14)

        wb.save(workbook_path)
        print(f"[+] Appended '{sheet_name}' sheet directly into: {workbook_path}")
        return workbook_path


