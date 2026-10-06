"""Live end-to-end check of the AOC-4 financial figure judge on a REAL OCR folder.

This is the figure-judge counterpart of ``run_live_llm_judge_check.py`` (common error
rules) and ``run_live_compliance_check.py`` (requirement rows). It reproduces the
production path for the 23 financial rows:

    OCR markdown files  ->  AOC4Parser.parse_all()      (real extraction)
                        ->  AOC4RuleEngine.evaluate_all() (real wiring)
                        ->  AOC4FigureLLMJudgeEngine.verify()  (Ollama model)
                        ->  F{row} verdict / G{row} evidence

The extracted figures in column B are never modified; the script reports the extracted
value, the model's verdict, the figure the document shows, and the evidence quote, so
accuracy can be audited row by row.

Usage
-----
    python automation_engine/modules/aoc4/tests/run_live_figure_check.py \
        "Y:\\Main-OCR-main\\Main-OCR-main\\data\\Macremedecor\\ocr_output"

Options
-------
    --url URL         Ollama base URL (default: AOC4_OLLAMA_URL or
                      DEFAULT_URL = http://192.168.112.2:11434)
    --model NAME      Model name (default: AOC4_LLM_MODEL or DEFAULT_MODEL = qwen3.5:4b)
    --timeout SEC     Per-request timeout (default: 300 - CPU inference is slow)
    --only TEXT       Only cross-check rows whose label contains TEXT
    --rows A,B,C      Only cross-check these sheet rows
    --dump-prompts D  Write the exact prompt sent per row (audit trail)
    --report FILE     JSON report path (default: figure_live_report.json)
    --dry-run         Build and dump prompts without calling the model

Exit code is 0 when the run completed. The report always records which rows were
judged by the model, which returned no verdict, and how long each call took.
"""

import argparse
import json
import os
import sys
import time

TESTS_DIR = os.path.dirname(os.path.abspath(__file__))
AOC4_DIR = os.path.dirname(TESTS_DIR)                                  # .../modules/aoc4
# aoc4 -> modules -> automation_engine -> repository root
PROJECT_ROOT = os.path.dirname(os.path.dirname(os.path.dirname(AOC4_DIR)))

for _path in (PROJECT_ROOT, AOC4_DIR):
    if _path not in sys.path:
        sys.path.insert(0, _path)

import openpyxl

from automation_engine.modules.aoc4.parser import AOC4Parser
from automation_engine.modules.aoc4.rule_engine import AOC4RuleEngine
from automation_engine.modules.aoc4.llm.figure_llm_judge import AOC4FigureLLMJudgeEngine
from automation_engine.modules.aoc4.llm.llm_judge import AOC4LLMJudgeEngine

CONFIG_PATH = os.path.join(AOC4_DIR, "rules_config.json")

PRIVATE_SHEET_CANDIDATES = [
    "Compliance sheet for private",
    "compliance for Private ",
    "compliance for Private",
]


class RecordingJudge:
    """Pass-through wrapper around the real figure judge that records each call.

    ``AOC4RuleEngine`` only ever calls ``is_judged_figure()``, ``name_for_figure()``
    and ``verify()``, so delegating those three keeps production behaviour identical
    while capturing latency, verdicts and errors for the report.
    """

    def __init__(self, inner, prompt_dump_dir=None, only=None, rows=None, dry_run=False,
                 row_map=None):
        self.inner = inner
        self.prompt_dump_dir = prompt_dump_dir
        self.only = only
        self.rows = set(rows or ())
        self.dry_run = dry_run
        #: figure key -> sheet row, so the report lines up with the Excel output
        #: (``verify()`` is keyed by the figure name, not the row).
        self.row_map = row_map or {}
        self.calls = []

    def is_judged_figure(self, label):
        return self.inner.is_judged_figure(label)

    def name_for_figure(self, label):
        return self.inner.name_for_figure(label)

    def verify(self, key, expected, input_data):
        spec = self.inner._FIGURES[key]
        row = self.row_map.get(key)

        # Row filters are applied here rather than in is_judged_figure() so the rows
        # that are skipped still fall through to the plain extracted figure, which is
        # what production does for a row it cannot ask about.
        if self.rows and row not in self.rows:
            return None
        if self.only and self.only.lower() not in spec["name"].lower():
            return None

        full_text = input_data.get("full_text") or ""

        # Rebuild the exact window/prompt the real judge will use.
        try:
            window = self.inner._relevant_window(full_text, rule_name=key)
            prompt = self.inner._build_figure_prompt(
                row_name=spec["name"], kind=spec["kind"],
                expected=expected, window=window)
        except Exception:  # pragma: no cover - defensive
            window, prompt = "", ""

        if self.prompt_dump_dir:
            os.makedirs(self.prompt_dump_dir, exist_ok=True)
            with open(os.path.join(self.prompt_dump_dir, "%s.prompt.txt" % key),
                      "w", encoding="utf-8") as fh:
                fh.write(prompt)

        started = time.time()
        result, error = None, None
        if self.dry_run:
            error = "dry-run"
        else:
            try:
                result = self.inner.verify(key, expected, input_data)
                if result is None:
                    error = "no verdict returned"
            except Exception as exc:
                error = "%s: %s" % (type(exc).__name__, exc)
        elapsed = time.time() - started

        self.calls.append({
            "key": key,
            "row": row,
            "name": spec["name"],
            "kind": spec["kind"],
            "expected": expected,
            "verdict": (result or {}).get("verdict"),
            "seen": (result or {}).get("seen"),
            "quote": (result or {}).get("quote"),
            "reason": (result or {}).get("reason"),
            "seconds": round(elapsed, 2),
            "window_chars": len(window),
            "document_chars": len(full_text),
            "error": error,
        })
        return result


def classify(name):
    """Map an OCR markdown filename onto the document role the parser expects.

    ``parse_all`` uses the ``financials`` key to detect the target year, and treats any
    key or filename containing prev/previous/prior as last year's document, so the
    role is encoded in the key rather than only the filename.
    """
    low = name.lower()
    if any(h in low for h in ("previous", "previouus", "prev ", "prevos", "_previous",
                              "last year", "old ", "23-24", "23_24", "fy 2023",
                              "2023-24")):
        return "previous_financials"
    if "masterdata" in low or "master data" in low:
        return "master_data"
    if "board" in low:
        return "board_report"
    if "audit" in low:
        return "audit_report"
    if any(k in low for k in ("financial", "balance", "ar and", "annual return")):
        return "financials"
    return "document"


def build_docs(folder):
    names = sorted(n for n in os.listdir(folder) if n.lower().endswith(".md"))
    if not names:
        raise SystemExit("No .md OCR files found in %s" % folder)
    docs = {}
    for name in names:
        docs[classify(name)] = os.path.join(folder, name)
    return docs, names


def health_check(url, timeout=5):
    import requests
    try:
        resp = requests.get("%s/api/tags" % url, timeout=timeout)
        resp.raise_for_status()
        return True, [m.get("name") for m in resp.json().get("models", [])]
    except Exception as exc:
        return False, "%s: %s" % (type(exc).__name__, exc)


def resolve_private_sheet(workbook):
    wb = openpyxl.load_workbook(workbook, data_only=True)
    for name in PRIVATE_SHEET_CANDIDATES:
        if name in wb.sheetnames:
            return name
    raise SystemExit("Private compliance sheet not found; sheets are: %s" % wb.sheetnames)


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("folder", nargs="?", default=None)
    parser.add_argument("--url", default=os.environ.get("AOC4_OLLAMA_URL",
                                                        AOC4LLMJudgeEngine.DEFAULT_URL))
    parser.add_argument("--model", default=os.environ.get("AOC4_LLM_MODEL",
                                                          AOC4LLMJudgeEngine.DEFAULT_MODEL))
    parser.add_argument("--timeout", type=float, default=300.0)
    parser.add_argument("--only", default=None)
    parser.add_argument("--rows", default=None,
                        help="comma separated sheet rows, e.g. 5,6,7")
    parser.add_argument("--dump-prompts", default=None)
    parser.add_argument("--report", default="figure_live_report.json")
    parser.add_argument("--dry-run", action="store_true")
    args = parser.parse_args(argv)

    folder = os.path.abspath(args.folder) if args.folder else None
    if not folder:
        raise SystemExit("Pass the path to an ocr_output folder.")
    if not os.path.isdir(folder):
        raise SystemExit("Input folder not found: %s" % folder)

    rows = [int(r) for r in args.rows.split(",")] if args.rows else None

    print("=" * 104)
    print("AOC-4 financial figure judge live check")
    print("  input folder : %s" % folder)
    print("   ollama url   : %s" % args.url)
    print("   model        : %s" % args.model)
    print("=" * 104)

    docs, names = build_docs(folder)
    print("\n[1/4] Document roles (parse_all keys):")
    for key, path in docs.items():
        print("  %-20s <- %s" % (key, os.path.basename(path)))

    ocr_parser = AOC4Parser(CONFIG_PATH)
    started = time.time()
    extracted = ocr_parser.parse_all(docs, {})
    print("  => full_text = %s chars, parse %.2fs, FS target year: %s" % (
        format(len(extracted.get("full_text", "")), ","), time.time() - started,
        extracted.get("financials_target_year")))

    engine = AOC4RuleEngine(CONFIG_PATH)
    sheet_name = resolve_private_sheet(engine.output_skeletal_path)
    sheet = openpyxl.load_workbook(engine.output_skeletal_path, data_only=True)[sheet_name]
    block = engine._financial_block_rows(sheet, 1)
    print("  => financial block on '%s': rows %s" % (sheet_name, block))

    reachable, detail = health_check(args.url)
    print("\n[2/4] Ollama %s - %s" % ("REACHABLE" if reachable else "UNREACHABLE", detail))

    inner = AOC4FigureLLMJudgeEngine(ollama_url=args.url, model=args.model,
                                     timeout=args.timeout)

    # Map each figure key back to the sheet row that holds it, mirroring how
    # _verify_financial_figures walks the block. ``Exprt`` appears twice, so a
    # repeated key keeps its first row here and is reported against that row.
    row_map = {}
    for row in block:
        key = inner.name_for_figure(sheet.cell(row=row, column=1).value)
        if key and key not in row_map:
            row_map[key] = row
    print("  => %d distinct figure(s) mapped to rows %s"
          % (len(row_map), sorted(row_map.values())))

    recorder = RecordingJudge(inner, prompt_dump_dir=args.dump_prompts,
                              only=args.only, rows=rows, dry_run=args.dry_run,
                              row_map=row_map)
    engine.figure_judge = recorder

    print("\n[3/4] Running evaluate_all() with the live figure judge...")
    started = time.time()
    target_cells = engine.evaluate_all(extracted)
    total_seconds = time.time() - started

    print("\n[4/4] Writing %s" % args.report)
    cells = target_cells.get(sheet_name, {})
    verdicts = {c: v for c, v in cells.items() if c.startswith("F")}

    print("\n" + "-" * 104)
    print("%-5s %-46s %16s %-12s %16s %7s" % (
        "ROW", "LINE ITEM", "EXTRACTED", "VERDICT", "DOC SHOWS", "SECONDS"))
    print("-" * 104)
    for call in recorder.calls:
        exp = call["expected"]
        seen = call["seen"]
        print("%-5s %-46s %16s %-12s %16s %7s" % (
            call.get("row", "-"), call["name"][:46],
            "-" if exp is None else format(exp, ",.2f") if isinstance(exp, (int, float)) else str(exp)[:16],
            call["verdict"] or "no verdict",
            "-" if seen is None else format(seen, ",.2f"),
            "%.2f" % call["seconds"]))
    print("-" * 104)
    print("evaluate_all() total: %.2fs across %d model call(s); %d row verdict(s) written"
          % (total_seconds, len(recorder.calls), len(verdicts)))

    judged = [c for c in recorder.calls if c["verdict"]]
    print("\nEvidence quotes:")
    for call in judged:
        if call["quote"]:
            print("\n  %s -> %s" % (call["name"], call["verdict"]))
            print("    %s" % call["quote"][:400])

    report = {
        "input_folder": folder,
        "documents": {k: os.path.basename(v) for k, v in docs.items()},
        "full_text_chars": len(extracted.get("full_text", "")),
        "financials_target_year": extracted.get("financials_target_year"),
        "private_sheet": sheet_name,
        "financial_block_rows": block,
        "ollama_url": args.url,
        "model": args.model,
        "ollama_reachable": reachable,
        "ollama_detail": str(detail),
        "dry_run": args.dry_run,
        "total_seconds": round(total_seconds, 2),
        "model_calls": recorder.calls,
        "verdict_cells": verdicts,
    }
    with open(args.report, "w", encoding="utf-8") as fh:
        json.dump(report, fh, indent=2, default=str)
    print("\nReport written: %s" % os.path.abspath(args.report))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
